"""
agent/attachment_node.py
Nó LangGraph: Extrai texto e itens de anexos (PDF, Word, imagens).
Usa pypdf, python-docx e a bridge de IA para cobertura total de formatos.
"""
import io
import os
from agent.config import get_llm_response
from graph.email_reader import get_email_attachments, decode_attachment

SUPPORTED_IMAGE_TYPES = {"image/jpeg", "image/png", "image/gif", "image/webp"}
SUPPORTED_PDF = {"application/pdf"}
SUPPORTED_WORD = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
}

def _extract_pdf_text(content: bytes) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        return text.strip()
    except Exception as e:
        return f"[Erro ao ler PDF: {e}]"

def _extract_word_text(content: bytes) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(content))
        full_text = []
        
        # 1. Extrai Parágrafos
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        
        # 2. Extrai Tabelas (MUITO COMUM em cotações)
        for table in doc.tables:
            for row in table.rows:
                # Pega o texto de cada célula, limpando espaços
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                # Remove duplicatas consecutivas (comum em células mescladas)
                unique_cells = []
                for c in row_cells:
                    if not unique_cells or c != unique_cells[-1]:
                        unique_cells.append(c)
                
                if unique_cells:
                    full_text.append(" | ".join(unique_cells))
                    
        return "\n".join(full_text).strip()
    except Exception as e:
        return f"[Erro ao ler Word: {e}]"

def _extract_image_text(content: bytes, content_type: str) -> str:
    """Usa a bridge de IA (Gemini Vision ou Groq Vision) para extrair texto de imagens."""
    try:
        image_data = {
            "mime_type": content_type,
            "data": content
        }
        
        prompt = (
            "Analise esta imagem. Se for um pedido de compra, lista de produtos ou cotacao, "
            "extraia todos os itens, quantidades e descricoes em texto. "
            "Se nao for uma solicitacao comercial, diga 'Nenhum pedido encontrado'."
        )
        return get_llm_response(prompt, image_data=image_data)
    except Exception as e:
        return f"[Erro ao processar imagem com IA: {e}]"

def process_attachments(state: dict) -> dict:
    """
    Nó LangGraph: baixa e processa todos os anexos do e-mail.
    Adiciona o texto extraído ao estado como attachment_text.
    """
    token = state.get("access_token", "")
    email_data = state.get("email_data", {})
    message_id = email_data.get("id", "")
    has_attachments = email_data.get("hasAttachments", False)

    if not has_attachments or not message_id:
        return {**state, "attachment_text": ""}

    try:
        # Verifica se há imagens embutidas no corpo (inline)
        body_raw = email_data.get("body", {}).get("content", "")
        has_cid = "cid:" in body_raw
        has_data_uri = "data:image" in body_raw
        if has_cid or has_data_uri:
            print(f"[DEBUG ANEXO] Detectadas imagens no corpo: cid={has_cid}, data={has_data_uri}", flush=True)

        attachments = get_email_attachments(token, message_id)
        print(f"\\n[DEBUG ANEXO] {len(attachments)} anexos encontrados no e-mail.", flush=True)
    except Exception as e:
        return {**state, "attachment_text": f"[Erro ao buscar anexos: {e}]"}

    extracted_texts = []

    for attachment in attachments:
        name = attachment.get("name", "arquivo")
        content_type = attachment.get("contentType", "").lower()
        is_inline = attachment.get("isInline", False)
        
        print(f"\\n[DEBUG ANEXO] Lendo arquivo: {name} ({content_type}) | isInline: {is_inline}", flush=True)

        try:
            content = decode_attachment(attachment)
        except Exception as e:
            err_msg = f"[Erro ao decodificar {name}: {e}]"
            print(f"  {err_msg}", flush=True)
            extracted_texts.append(err_msg)
            continue

        text = ""
        if content_type in SUPPORTED_PDF:
            text = _extract_pdf_text(content)
            extracted_texts.append(f"📄 PDF [{name}]:\\n{text}")
        elif content_type in SUPPORTED_WORD:
            text = _extract_word_text(content)
            extracted_texts.append(f"📝 Word [{name}]:\\n{text}")
        elif content_type in SUPPORTED_IMAGE_TYPES:
            text = _extract_image_text(content, content_type)
            extracted_texts.append(f"🖼️ Imagem [{name}] (via IA):\\n{text}")
        else:
            text = f"⚠️ Formato não suportado: {name} ({content_type})"
            extracted_texts.append(text)
        
        print(f"  Conteudo extraido (parcial): {str(text)[:200]}...", flush=True)

    attachment_text = "\\n\\n".join(extracted_texts)
    return {**state, "attachment_text": attachment_text}
