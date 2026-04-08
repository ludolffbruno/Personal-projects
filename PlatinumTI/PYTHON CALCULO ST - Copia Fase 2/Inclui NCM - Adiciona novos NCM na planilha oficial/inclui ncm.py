import pandas as pd
from openpyxl import load_workbook

# Passo 1: Carregar a planilha "Incluir.xlsx" com pandas
incluir_df = pd.read_excel('Incluir.xlsx', sheet_name=0)

# Passo 2: Carregar a planilha "CALCULO_ST_DIFAL_COMPLETA.xlsx" com openpyxl
wb = load_workbook('CALCULO_ST_DIFAL_COMPLETA.xlsx')
ws = wb.active  # Assume que estamos trabalhando na aba ativa

# Passo 3: Definir as colunas (ajuste os índices conforme sua planilha)
ncm_col = 1  # Coluna A (NCM)
tem_reducao_col = 14  # Coluna N (tem reducao)

# Passo 4: Mapear os NCMs existentes para suas linhas
ncm_to_row = {}
for row in range(2, ws.max_row + 1):  # Começa na linha 2, assumindo cabeçalho na linha 1
    ncm = ws.cell(row=row, column=ncm_col).value
    if ncm:
        ncm_to_row[str(ncm)] = row

# Passo 5: Atualizar "tem reducao" para NCMs existentes
for ncm in incluir_df['NCM'].astype(str):
    if ncm in ncm_to_row:
        row = ncm_to_row[ncm]
        ws.cell(row=row, column=tem_reducao_col).value = 'SIM'

# Passo 6: Adicionar NCMs novos ao final da planilha
for index, row in incluir_df.iterrows():
    ncm = str(row['NCM'])
    if ncm not in ncm_to_row:
        # Adiciona nova linha com NCM e Produto
        nova_linha = ws.max_row + 1
        ws.cell(row=nova_linha, column=ncm_col).value = ncm
        ws.cell(row=nova_linha, column=2).value = row['PRODUTO']  # Assume Produto na coluna B

# Passo 7: Salvar a planilha com a formatação preservada
wb.save('CALCULO_ST_DIFAL_COMPLETA.xlsx')

print("Atualização concluída! A formatação original foi mantida.")
